"""Reportes por periodo y exportación a Excel."""

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from arbitros.config import ORGANIZACION
from arbitros.db.models import ESTADOS_PARTIDO, FORMAS_PAGO, POSICIONES
from arbitros.services.cedulas import NOMBRES_ESTADO, PENDIENTE, CedulaService
from arbitros.services.cobros import CobroService
from arbitros.services.pagos import PagoService
from arbitros.services.texto import sin_acentos


@dataclass
class FilaArbitraje:
    """Partidos arbitrados por un árbitro en el periodo."""

    nombre: str
    por_posicion: dict[str, int] = field(default_factory=dict)
    total_partidos: int = 0
    ganado: int = 0
    pendiente: int = 0


class ReporteService:
    def __init__(self, pagos: PagoService, cedulas: CedulaService, cobros: CobroService):
        self.pagos = pagos
        self.cedulas = cedulas
        self.cobros = cobros

    def arbitrajes(self, desde: date | None, hasta: date | None, liga_id: int | None = None) -> list[FilaArbitraje]:
        filas: dict[int, FilaArbitraje] = {}
        for l in self.pagos.lineas(desde, hasta, liga_id):
            if not l.partido.se_juega:
                continue
            f = filas.setdefault(l.arbitro.id, FilaArbitraje(l.arbitro.nombre))
            f.por_posicion[l.posicion] = f.por_posicion.get(l.posicion, 0) + 1
            f.total_partidos += 1
            f.ganado += l.tarifa
            if l.jugado:
                f.pendiente += l.saldo
        return sorted(filas.values(), key=lambda f: (-f.total_partidos, sin_acentos(f.nombre)))

    def partidos_por_estado(self, desde: date | None, hasta: date | None, liga_id: int | None = None) -> dict[str, int]:
        conteo = {e: 0 for e in ESTADOS_PARTIDO}
        # CedulaService.listar ya recorre todos los partidos del periodo con su liga
        for fila in self.cedulas.listar(liga_id=liga_id, desde=desde, hasta=hasta):
            conteo[fila.partido.estado] = conteo.get(fila.partido.estado, 0) + 1
        return conteo

    def exportar_excel(self, ruta: Path, desde: date | None, hasta: date | None, liga_id: int | None = None) -> Path:
        libro = Workbook()
        periodo = f"Del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}" if desde and hasta else "Todo el historial"

        hoja = libro.active
        hoja.title = "Resumen por árbitro"
        filas = []
        for r in self.pagos.resumen_por_arbitro(desde, hasta, liga_id):
            filas.append([r.arbitro.nombre, r.partidos, _p(r.total), _p(r.pagado_campo), _p(r.liquidado),
                          _p(r.pendiente), _p(r.por_jugar)])
        _tabla(hoja, periodo, ["Árbitro", "Partidos", "Total", "Pagado en campo", "Pagado después",
                                "Pendiente", "Por jugarse"], filas, moneda={3, 4, 5, 6, 7})

        hoja = libro.create_sheet("Arbitrajes")
        posiciones = list(POSICIONES)
        filas = [[f.nombre, *[f.por_posicion.get(p, 0) for p in posiciones], f.total_partidos, _p(f.ganado)]
                 for f in self.arbitrajes(desde, hasta, liga_id)]
        _tabla(hoja, periodo, ["Árbitro", *POSICIONES.values(), "Total", "Ganado"], filas,
               moneda={len(posiciones) + 3})

        hoja = libro.create_sheet("Detalle de pagos")
        filas = []
        for l in self.pagos.lineas(desde, hasta, liga_id):
            p = l.partido
            filas.append([p.fecha, f"{p.hora:%H:%M}", l.liga.nombre, l.jornada,
                          p.categoria.nombre if p.categoria else "", p.campo.nombre if p.campo else "",
                          f"{p.local.nombre} vs {p.visitante.nombre}", ESTADOS_PARTIDO.get(p.estado, p.estado),
                          l.arbitro.nombre, POSICIONES[l.posicion], _p(l.tarifa), _p(l.pagado_campo),
                          _p(l.pagado_despues), _p(l.saldo)])
        _tabla(hoja, periodo, ["Fecha", "Hora", "Liga", "Jornada", "Categoría", "Campo", "Partido", "Estado",
                                "Árbitro", "Posición", "Tarifa", "Pagado en campo", "Pagado después", "Saldo"],
               filas, moneda={11, 12, 13, 14}, fechas={1})

        hoja = libro.create_sheet("Historial de pagos")
        filas = [[pago.fecha, pago.arbitro.nombre, _p(pago.monto_centavos), n, pago.nota or ""]
                 for pago, n in self.pagos.historial()
                 if (not desde or pago.fecha >= desde) and (not hasta or pago.fecha <= hasta)]
        _tabla(hoja, periodo, ["Fecha", "Árbitro", "Monto", "Partidos", "Nota"], filas, moneda={3}, fechas={1})

        hoja = libro.create_sheet("Pagos de equipos")
        filas = [[r.liga.nombre, r.equipo.nombre, r.partidos, _p(r.total), _p(r.pagado), _p(r.pendiente),
                  _p(r.por_jugar), _p(r.a_favor)]
                 for r in self.cobros.resumen_por_equipo(desde, hasta, liga_id)]
        _tabla(hoja, periodo, ["Liga", "Equipo", "Partidos", "Arbitrajes", "Pagado", "Debe", "Por jugarse",
                                "Saldo a favor"], filas, moneda={4, 5, 6, 7, 8})

        hoja = libro.create_sheet("Detalle de equipos")
        filas = []
        for l in self.cobros.lineas(desde, hasta, liga_id):
            p = l.partido
            filas.append([p.fecha, f"{p.hora:%H:%M}", l.liga.nombre, l.jornada,
                          p.categoria.nombre if p.categoria else "", l.equipo.nombre, l.rival.nombre,
                          _p(l.cuota), _p(l.pagado), _p(l.abonado), _p(l.saldo), l.fecha_pago or "",
                          FORMAS_PAGO.get(l.forma_pago, "") if l.forma_pago else "", l.nota or ""])
        _tabla(hoja, periodo, ["Fecha", "Hora", "Liga", "Jornada", "Categoría", "Equipo", "Rival",
                                "Arbitraje", "Pagado", "Pago adelantado", "Saldo", "Fecha de pago", "Forma de pago",
                                "Nota"],
               filas, moneda={8, 9, 10, 11}, fechas={1, 12})

        hoja = libro.create_sheet("Pagos adelantados")
        filas = [[a.fecha, a.equipo.nombre, _p(a.monto_centavos), FORMAS_PAGO.get(a.forma_pago, "") if a.forma_pago else "",
                  a.nota or ""]
                 for a in self.cobros.abonos(liga_id=liga_id)
                 if (not desde or a.fecha >= desde) and (not hasta or a.fecha <= hasta)]
        _tabla(hoja, periodo, ["Fecha", "Equipo", "Monto", "Forma de pago", "Nota"], filas, moneda={3}, fechas={1})

        hoja = libro.create_sheet("Cédulas")
        filas = [[c.partido.fecha, f"{c.partido.hora:%H:%M}", c.liga.nombre, c.jornada,
                  f"{c.partido.local.nombre} vs {c.partido.visitante.nombre}", NOMBRES_ESTADO[c.estado],
                  len(c.archivos), c.motivo_excepcion or ""]
                 for c in self.cedulas.listar(liga_id=liga_id, desde=desde, hasta=hasta)]
        _tabla(hoja, periodo, ["Fecha", "Hora", "Liga", "Jornada", "Partido", "Estado", "Archivos", "Excepción"],
               filas, fechas={1})

        ruta = ruta if ruta.suffix.lower() == ".xlsx" else ruta.with_suffix(".xlsx")
        libro.save(ruta)
        return ruta

    def cedulas_pendientes(self, desde: date | None, hasta: date | None, liga_id: int | None = None):
        return self.cedulas.listar(liga_id=liga_id, desde=desde, hasta=hasta, estados={PENDIENTE})


def _p(centavos: int) -> float:
    return round(centavos / 100, 2)


ROJO = "D41014"


def _tabla(hoja, periodo: str, encabezados: list[str], filas: list[list], moneda: set[int] = frozenset(),
           fechas: set[int] = frozenset()) -> None:
    hoja.append([ORGANIZACION])
    hoja.append([f"{hoja.title} · {periodo}"])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja["A2"].font = Font(italic=True, color="666666")
    hoja.append([])
    hoja.append(encabezados)
    for celda in hoja[4]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=ROJO)
        celda.alignment = Alignment(horizontal="center")
    for fila in filas:
        hoja.append(fila)
    for col in range(1, len(encabezados) + 1):
        letra = get_column_letter(col)
        for celda in hoja[letra][4:]:
            if col in moneda:
                celda.number_format = '"$"#,##0.00'
            if col in fechas:
                celda.number_format = "dd/mm/yyyy"
        ancho = max([len(str(encabezados[col - 1]))] + [len(str(f[col - 1])) for f in filas if len(f) >= col])
        hoja.column_dimensions[letra].width = min(max(ancho + 2, 10), 45)
    hoja.freeze_panes = "A5"
